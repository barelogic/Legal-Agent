"""Pure render helpers: verifiability rules enforced BEFORE Streamlit.

Grounding rule: a claim is renderable as a fact ONLY if it carries at
least one ``chunk_id`` AND a non-empty verbatim ``quote``. Anything else
raises and must be shown as dropped/refused, never as a fact.
"""

from __future__ import annotations

import difflib
import html
import json
import re
from pathlib import Path

from contracts.schemas import Answer, Citation, Claim, Contradiction, Doc

MARKER_RE = re.compile(r"\[(c\d+)\]")

STATUS_BADGE: dict[str, str] = {
    "verified":    "✅ verified",
    "unsupported": "⚠️ unsupported",
    "removed":     "⛔ removed",
    "unverified":  "❔ unverified",
}

# Inline HTML badges (used in tables and chips only — not in prose)
_BADGE_HTML: dict[str, str] = {
    "verified":    '<span class="badge badge-v">✅ Verified</span>',
    "unsupported": '<span class="badge badge-u">⚠️ Unsupported</span>',
    "removed":     '<span class="badge badge-r">⛔ Removed</span>',
    "unverified":  '<span class="badge badge-n">❔ Unverified</span>',
}

# Kept for backward compat — app.py imports USER_VALUE_CSS (now empty string)
USER_VALUE_CSS = ""


# ─── Doc / chunk helpers ─────────────────────────────────────────────────────

def docs_by_id(docs: list[Doc]) -> dict[str, Doc]:
    """Index Docs by doc_id for source resolution."""
    return {d.doc_id: d for d in docs}


def doc_id_of_chunk(chunk_id: str) -> str:
    """Extract doc_id from a ``{doc_id}::...`` chunk id (fallback '')."""
    return chunk_id.split("::")[0] if "::" in chunk_id else ""


def validate_claim_renderable(claim: Claim) -> None:
    """Raise ValueError if the claim must not be shown as a verified fact."""
    if not claim.chunk_ids:
        raise ValueError(f"claim {claim.claim_id} has no chunk_ids: refusing to render")
    if not claim.quote.strip():
        raise ValueError(f"claim {claim.claim_id} has no quote: refusing to render")


def source_label(chunk_id: str, docs: dict[str, Doc]) -> str:
    """Human-readable source label, honest when Doc metadata is missing."""
    doc = docs.get(doc_id_of_chunk(chunk_id))
    if doc is None:
        return f"{chunk_id} (unregistered doc)"
    bits = [doc.title]
    if doc.citation:
        bits.append(doc.citation)
    return f"{' — '.join(bits)} · `{chunk_id}`"


# ─── Claim helpers ───────────────────────────────────────────────────────────

def claim_to_markdown(claim: Claim, docs: dict[str, Doc]) -> str:
    """Render ONE verified claim with its quote + source links. No source → raise."""
    validate_claim_renderable(claim)
    lines = [claim.text, ""]
    lines.append(f"> {claim.quote}")
    lines.append("")
    for cid in claim.chunk_ids:
        lines.append(f"- Source: {source_label(cid, docs)}")
        doc = docs.get(doc_id_of_chunk(cid))
        if doc is not None and doc.source_url:
            lines.append(f"  <{doc.source_url}|open source>")
        else:
            lines.append("  _local doc — no external link_")
    if claim.verifier_note:
        lines.append(f"\n_Verifier: {claim.verifier_note}_")
    return "\n".join(lines)


def verified_claims(answer: Answer) -> list[Claim]:
    """Claims safe to show as facts (verified + renderable)."""
    out = []
    for c in answer.claims:
        if c.status != "verified":
            continue
        try:
            validate_claim_renderable(c)
        except ValueError:
            continue
        out.append(c)
    return out


def dropped_claims(answer: Answer) -> list[Claim]:
    """Everything else: audit trail, shown collapsed — never as facts."""
    v_ids = {c.claim_id for c in verified_claims(answer)}
    return [c for c in answer.claims if c.claim_id not in v_ids]


def split_text_markers(text: str) -> list[tuple[str, str | None]]:
    """Split ``Answer.text`` into (segment, claim_id|None) parts on [cN]."""
    parts: list[tuple[str, str | None]] = []
    last = 0
    for m in MARKER_RE.finditer(text):
        if m.start() > last:
            parts.append((text[last:m.start()], None))
        parts.append((m.group(0), m.group(1)))
        last = m.end()
    if last < len(text):
        parts.append((text[last:], None))
    if not parts:
        parts.append((text, None))
    return parts


def claim_markers(text: str) -> list[str]:
    """Ordered unique claim ids referenced as [cN] in text."""
    seen: list[str] = []
    for m in MARKER_RE.finditer(text):
        if m.group(1) not in seen:
            seen.append(m.group(1))
    return seen


def highlight_quote(chunk_text: str, quote: str) -> str:
    """Escape HTML and wrap the first verbatim quote occurrence in <mark>."""
    safe = html.escape(chunk_text)
    q = html.escape(quote)
    if quote and quote in chunk_text and q in safe:
        return safe.replace(q, f"<mark>{q}</mark>", 1)
    return safe


def status_badge(status: str) -> str:
    """Badge label for a claim status; unknown statuses pass through."""
    return STATUS_BADGE.get(status, status)


def status_badge_html(status: str) -> str:
    """Inline HTML badge for use in tables / panels."""
    return _BADGE_HTML.get(
        status,
        f'<span class="badge badge-n">{html.escape(status)}</span>',
    )


# ─── Citations ───────────────────────────────────────────────────────────────

def citation_to_markdown(cite: Citation, docs: dict[str, Doc]) -> str:
    """Render one citation; unresolved citations are flagged, never linked."""
    if not cite.resolved or not cite.doc_id or cite.doc_id not in docs:
        return f"`{cite.raw}` → ⚠️ unresolved (not in registry)"
    doc = docs[cite.doc_id]
    if doc.source_url:
        return f"`{cite.raw}` → [{doc.title}]({doc.source_url})"
    cited = f" ({doc.citation})" if doc.citation else ""
    return f"`{cite.raw}` → {doc.title}{cited} · `{doc.doc_id}` (local doc)"


# ─── Verifier summary ────────────────────────────────────────────────────────

def verifier_summary(answer: Answer) -> dict:
    """Verifier numbers for the summary bar, with provenance labels."""
    good = verified_claims(answer)
    bad = dropped_claims(answer)
    trace = answer.trace if isinstance(answer.trace, dict) else {}
    traced = trace.get("dropped")
    if isinstance(traced, int) and traced >= 0:
        dropped_n, source = traced, "trace"
    else:
        dropped_n, source = len(bad), "claims"
    trace_reasons = trace.get("dropped_reasons")
    if isinstance(trace_reasons, list) and trace_reasons:
        reasons = [str(r) for r in trace_reasons]
        reasons_source = "backend trace"
    else:
        reasons = [
            f"`{c.claim_id}` [{c.status}]"
            + (f" — {c.verifier_note}" if c.verifier_note else "")
            for c in bad
        ]
        reasons_source = "claim status/verifier notes (trace has no per-claim reasons)"
    return {
        "verified_n": len(good),
        "dropped_n": dropped_n,
        "dropped_source": source,
        "dropped_reasons": reasons,
        "reasons_source": reasons_source,
        "confidence": answer.confidence,
    }


def verifier_bar_html(summary: dict) -> str:
    """Verifier row: counts + colour-coded confidence gauge."""
    conf = summary["confidence"]
    if conf is not None:
        pct = int(conf * 100)
        colour = "#16A34A" if pct >= 70 else "#F59E0B" if pct >= 40 else "#DC2626"
        gauge = (
            f'<div class="gauge-wrap">'
            f'<div class="gauge-track">'
            f'<div class="gauge-fill" style="width:{pct}%;background:{colour}"></div>'
            f'</div>'
            f'<span class="gauge-pct" style="color:{colour}">{pct}%</span>'
            f'</div>'
        )
    else:
        gauge = '<span style="color:#9CA3AF;font-size:0.78rem">not provided</span>'

    return (
        f'<div style="display:flex;align-items:center;gap:24px;'
        f'padding:10px 16px;background:#F9FAFB;border:1px solid #E5E7EB;'
        f'border-radius:8px;font-size:0.82rem;color:#374151">'
        f'<span>✅ <strong>{summary["verified_n"]}</strong> verified</span>'
        f'<span style="color:#E5E7EB">|</span>'
        f'<span>⛔ <strong>{summary["dropped_n"]}</strong> dropped</span>'
        f'<span style="color:#E5E7EB">|</span>'
        f'<span style="flex:1;display:flex;align-items:center;gap:10px">'
        f'<span style="color:#6B7280">Confidence</span>{gauge}</span>'
        f'</div>'
    )


# ─── User-provided / missing chips ───────────────────────────────────────────

def user_value_html(value: str) -> str:
    """Wrap a user-supplied value in a visually distinct amber badge."""
    return (
        '<span class="user-provided"><span class="user-provided-tag">user-provided</span>'
        f"{html.escape(value)}</span>"
    )


def user_provided_chip(value: str) -> str:
    """Inline [USER-PROVIDED: value] marker in amber style."""
    return (
        '<span class="user-provided"><span class="user-provided-tag">user-provided</span>'
        f"[USER-PROVIDED: {html.escape(value)}]</span>"
    )


def missing_chip(field: str) -> str:
    """Inline red [MISSING: field] placeholder."""
    return f'<span class="missing-chip">[MISSING: {html.escape(field)}]</span>'


def copy_button_html(text: str, label: str = "📋 Copy") -> str:
    """Clipboard copy button."""
    safe = html.escape(text, quote=True).replace("'", "&#39;")
    return f"<button class='clip-btn' onclick=\"_clip('{safe}',this)\">{label}</button>"


# ─── Templates + precheck ────────────────────────────────────────────────────

TEMPLATE_IDS = ("bail_application", "legal_notice", "affidavit")


def _templates_dir() -> Path:
    return Path(__file__).resolve().parent.parent / "templates"


def load_template(template_id: str) -> dict:
    """Load one UI-owned template JSON; raises ValueError if unknown."""
    if template_id not in TEMPLATE_IDS:
        raise ValueError(f"unknown template {template_id!r}")
    return json.loads((_templates_dir() / f"{template_id}.json").read_text(encoding="utf-8"))


def required_fields(template: dict) -> list[dict]:
    """Required field dicts of a template JSON (in file order)."""
    return [f for f in template.get("required_fields", []) if f.get("required")]


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _field_flagged(norm_name: str, missing: list) -> object | None:
    """Return the MissingInfo entry flagging this template field, if any."""
    for m in missing:
        nm = _norm(m.field)
        if not nm or not norm_name:
            continue
        if nm == norm_name or nm in norm_name or norm_name in nm:
            return m
    return None


def precheck_summary(answer: Answer, template: dict) -> dict:
    """Draft readiness derived from trace.field_status (backend) or client-side fallback."""
    fields = required_fields(template)
    rows: list[dict] = []
    
    trace = answer.trace if isinstance(answer.trace, dict) else {}
    field_status = trace.get("field_status")
    
    if isinstance(field_status, dict):
        # Backend-provided field_status is available
        for f in fields:
            status = field_status.get(f["name"], "missing")
            # find matching missing_info for why_needed
            m = next((m for m in answer.missing_info if m.field == f["name"]), None)
            rows.append({
                "name": f["name"],
                "description": f.get("description", ""),
                "status": status,
                "why_needed": m.why_needed if m is not None else "",
                "searched_in": list(m.searched_in) if m is not None else [],
            })
        unconfirmed = False
        sourced_n = sum(1 for r in rows if r["status"] in ("sourced", "user_provided"))
    else:
        # Client-side fallback derivation
        for f in fields:
            m = _field_flagged(_norm(f["name"]), answer.missing_info)
            rows.append({
                "name": f["name"],
                "description": f.get("description", ""),
                "status": "missing" if m is not None else "sourced",
                "why_needed": m.why_needed if m is not None else "",
                "searched_in": list(m.searched_in) if m is not None else [],
            })
        unconfirmed = not answer.missing_info
        sourced_n = sum(1 for r in rows if r["status"] == "sourced")
        
    return {
        "template_id": template.get("template_id", ""),
        "total": len(rows),
        "sourced_n": sourced_n,
        "missing_n": len(rows) - sourced_n,
        "unconfirmed": unconfirmed,
        "rows": rows,
    }


def confidence_text(answer: Answer) -> str:
    """Honest one-line confidence readout (backend often sends None)."""
    if answer.confidence is None:
        return "confidence: not provided by backend"
    return f"confidence: {answer.confidence:.2f}"


# ─── Draft HTML ──────────────────────────────────────────────────────────────

def _claims_by_id(answer: Answer) -> dict[str, Claim]:
    return {c.claim_id: c for c in answer.claims}


def draft_html(answer: Answer) -> str:
    """Render Answer.text as HTML: sourced runs underlined w/ quote tooltip."""
    by_id = _claims_by_id(answer)
    out: list[str] = []
    buf = ""
    missing_re = re.compile(r"\[MISSING: (.*?)\]")
    user_re = re.compile(r"\[USER-PROVIDED: (.*?)\]")

    def _style_literal_markers(escaped_text: str) -> str:
        # Replaces literal markers with their corresponding styled chips.
        # escaped_text has already passed through html.escape
        text = missing_re.sub(lambda m: missing_chip(html.unescape(m.group(1))), escaped_text)
        text = user_re.sub(lambda m: user_provided_chip(html.unescape(m.group(1))), text)
        return text

    def flush(marker_cid: str | None) -> None:
        nonlocal buf
        if not buf:
            return
        claim = by_id.get(marker_cid) if marker_cid else None
        quote = claim.quote if claim is not None else ""
        try:
            if claim is not None:
                validate_claim_renderable(claim)
            else:
                raise ValueError("no claim")
            styled_buf = _style_literal_markers(html.escape(buf))
            out.append(
                f'<span class="sourced" title="{html.escape(quote)}">'
                f"{styled_buf}</span>"
            )
        except ValueError:
            out.append(_style_literal_markers(html.escape(buf)))
        buf = ""

    for seg, cid in split_text_markers(answer.text):
        if cid is None:
            buf += seg
        else:
            flush(cid)
            out.append(f" <strong>[{cid}]</strong>")
    flush(None)
    return "".join(out)


def still_needed_html(answer: Answer, provided_values: dict[str, str]) -> str:
    """Red placeholders / amber user-values for missing fields."""
    bits: list[str] = []
    real_missing = [m for m in answer.missing_info if m.field != "exhaustive coverage"]
    for m in real_missing:
        val = (provided_values or {}).get(m.field, "").strip()
        if val:
            bits.append(f"{missing_chip(m.field)} → {user_provided_chip(val)}")
        else:
            bits.append(missing_chip(m.field))
    return "<br>".join(bits)


# ─── DOCX export ─────────────────────────────────────────────────────────────

def draft_segments(answer: Answer, provided_values: dict[str, str]) -> dict:
    """Segment the draft body for screen-parity rendering."""
    body: list[list[tuple[str, str]]] = []
    for line in answer.text.splitlines():
        if not line.strip():
            continue
        para: list[tuple[str, str]] = []
        for seg, cid in split_text_markers(line):
            if cid is None:
                if seg:
                    para.append((seg, "text"))
            else:
                para.append((f"[{cid}]", "marker"))
        if para:
            body.append(para)
    still: list[list[tuple[str, str]]] = []
    for m in answer.missing_info:
        val = (provided_values or {}).get(m.field, "").strip()
        if val:
            still.append([
                (f"[MISSING: {m.field}]", "missing"),
                (" → ", "text"),
                (f"[USER-PROVIDED: {val}]", "user"),
            ])
        else:
            still.append([(f"[MISSING: {m.field}]", "missing")])
    return {"body": body, "still_needed": still}


def export_draft_docx(
    answer: Answer,
    question: str,
    provided_values: dict[str, str],
    docs: dict[str, Doc],
) -> bytes:
    """Build the draft DOCX: question, draft, still-needed, Sources appendix."""
    try:
        from docx import Document
        from docx.enum.text import WD_COLOR_INDEX
    except ImportError as e:
        raise RuntimeError(
            "python-docx is not installed (pip install -r ui/requirements.txt)"
        ) from e

    doc = Document()
    doc.add_heading("Draft", level=1)
    if question:
        p = doc.add_paragraph()
        p.add_run("Question (user-provided): ").italic = True
        p.add_run(question)

    segs = draft_segments(answer, provided_values)
    for para in segs["body"]:
        p = doc.add_paragraph()
        for text, kind in para:
            run = p.add_run(text)
            if kind == "marker":
                run.bold = True

    if segs["still_needed"]:
        doc.add_heading("Still needed", level=1)
        for para in segs["still_needed"]:
            p = doc.add_paragraph()
            for text, kind in para:
                run = p.add_run(text)
                run.bold = True
                if kind == "missing":
                    run.font.highlight_color = WD_COLOR_INDEX.RED
                elif kind == "user":
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_heading("Sources", level=1)
    for c in answer.claims:
        p = doc.add_paragraph()
        p.add_run(f"[{c.claim_id}] {status_badge(c.status)}").bold = True
        p.add_run(f" {c.text}")
        if not c.chunk_ids or not c.quote.strip():
            doc.add_paragraph("(no verified quote — not a sourced fact)").italic = True
            continue
        for cid in c.chunk_ids:
            d = docs.get(doc_id_of_chunk(cid))
            title = d.title if d else "(unregistered doc)"
            cite = f" ({d.citation})" if d and d.citation else ""
            doc.add_paragraph(f"{cid} — {title}{cite}")
            doc.add_paragraph(f"\u201c{c.quote}\u201d")

    from io import BytesIO
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── Contradiction diff ───────────────────────────────────────────────────────

def diff_sides(a: str, b: str) -> tuple[str, str]:
    """Word-level diff of two chunk texts; differing tokens get <mark>."""
    ta, tb = a.split(), b.split()
    sm = difflib.SequenceMatcher(None, ta, tb, autojunk=False)
    ha, hb = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            ha.extend(html.escape(t) for t in ta[i1:i2])
            hb.extend(html.escape(t) for t in tb[j1:j2])
        else:
            ha.extend(f"<mark>{html.escape(t)}</mark>" for t in ta[i1:i2])
            hb.extend(f"<mark>{html.escape(t)}</mark>" for t in tb[j1:j2])
    return " ".join(ha), " ".join(hb)


def contradiction_view_model(contra: Contradiction, fetched: dict[str, object | None]) -> dict:
    """View model for one contradiction: texts (or None when unfetched)."""
    return {
        "description": contra.description,
        "chunk_a": contra.claim_a,
        "chunk_b": contra.claim_b,
        "text_a": getattr(fetched.get(contra.claim_a), "text", None),
        "text_b": getattr(fetched.get(contra.claim_b), "text", None),
    }
