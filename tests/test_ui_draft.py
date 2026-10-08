"""Draft-flow helpers: precheck, draft view, verifier bar, contradictions.

Fixtures-first: every helper is exercised against fixture Answers shaped
like the frozen schema, including backend-empty fallbacks (no missing_info,
no trace keys, confidence None, failed chunk fetch).
"""

from contracts.schemas import Answer, Doc
from ui.render import (
    TEMPLATE_IDS,
    confidence_text,
    contradiction_view_model,
    diff_sides,
    docs_by_id,
    draft_html,
    draft_segments,
    export_draft_docx,
    load_template,
    missing_chip,
    precheck_summary,
    required_fields,
    still_needed_html,
    user_provided_chip,
    verifier_bar_html,
    verifier_summary,
)


def _answer(**kw) -> Answer:
    base: dict = {
        "workflow": "draft",
        "text": "- Bail rests on judicial discretion [c1].\n- FIR dated 1 Jan [c2].",
        "claims": [
            {
                "claim_id": "c1",
                "text": "Bail rests on judicial discretion.",
                "chunk_ids": ["bnss_2023::p1::c1"],
                "quote": "judicial discretion",
                "status": "verified",
            },
            {
                "claim_id": "c2",
                "text": "FIR dated 1 Jan.",
                "chunk_ids": ["case_file_demo::p1::c1"],
                "quote": "dated 2 Feb",
                "status": "unsupported",
                "verifier_note": "quote not verbatim in chunk",
            },
        ],
        "citations": [],
        "missing_info": [
            {
                "field": "FIR number",
                "why_needed": "links prayer to the case record",
                "searched_in": ["case_file_demo"],
            },
            {
                "field": "bail provision",
                "why_needed": "grounds need a statutory basis",
                "searched_in": ["bnss_2023"],
            },
        ],
        "contradictions": [
            {
                "description": "custody length differs",
                "claim_a": "case_file_demo::p1::c9",
                "claim_b": "case_file_demo::p2::c3",
            }
        ],
        "confidence": 0.82,
        "trace": {"retrieved_chunk_ids": ["bnss_2023::p1::c1"], "dropped": 1},
    }
    base.update(kw)
    return Answer(**base)


class _Chunk:
    def __init__(self, text: str):
        self.text = text


# --- templates -----------------------------------------------------------

def test_all_templates_load_with_required_fields():
    for tid in TEMPLATE_IDS:
        t = load_template(tid)
        assert t["template_id"] == tid
        assert required_fields(t), tid


def test_load_template_unknown_raises():
    try:
        load_template("nope")
    except ValueError:
        return
    raise AssertionError("expected ValueError")


# --- precheck --------------------------------------------------------------

def test_precheck_counts_flagged_fields():
    s = precheck_summary(_answer(), load_template("bail_application"))
    assert s["total"] == 19
    assert s["missing_n"] == 2  # FIR number + bail provision flagged
    assert s["sourced_n"] == 17
    assert s["unconfirmed"] is False
    flagged = {r["name"] for r in s["rows"] if r["status"] == "missing"}
    assert flagged == {"fir_number", "bail_provision"}
    fir = next(r for r in s["rows"] if r["name"] == "fir_number")
    assert fir["why_needed"] == "links prayer to the case record"
    assert fir["searched_in"] == ["case_file_demo"]


def test_precheck_unconfirmed_without_missing_list():
    s = precheck_summary(_answer(missing_info=[]), load_template("bail_application"))
    assert s["unconfirmed"] is True
    assert s["sourced_n"] == s["total"]  # nothing flagged — caller must caveat


def test_confidence_text_honest():
    assert "0.82" in confidence_text(_answer())
    assert "not provided" in confidence_text(_answer(confidence=None))


# --- draft view ------------------------------------------------------------

def test_draft_html_underlines_sourced_with_quote_tooltip():
    out = draft_html(_answer())
    assert 'class="sourced"' in out
    assert 'title="judicial discretion"' in out
    assert "[c1]" in out  # marker kept for the chip buttons


def test_draft_html_never_underlines_without_quote():
    out = draft_html(_answer(missing_info=[]))
    # c2 is unsupported but still carries chunk_ids+quote → underlined too
    assert out.count('class="sourced"') == 2
    bare = _answer(text="No markers here.", claims=[])
    plain = draft_html(bare)
    assert 'class="sourced"' not in plain
    assert "No markers here." in plain


def test_still_needed_placeholders_and_user_values():
    need = still_needed_html(_answer(), {})
    assert "[MISSING: FIR number]" in need
    assert "[MISSING: bail provision]" in need
    assert "missing-chip" in need
    mixed = still_needed_html(_answer(), {"FIR number": "0123/2024"})
    assert "[USER-PROVIDED: 0123/2024]" in mixed
    assert "user-provided" in mixed
    assert "[MISSING: bail provision]" in mixed  # still unsupplied
    assert still_needed_html(_answer(missing_info=[]), {}) == ""


def test_chips_labelled_and_escaped():
    assert "user-provided" in user_provided_chip("<x>")
    assert "<x>" not in user_provided_chip("<x>")
    assert "[MISSING:" in missing_chip("a<b>")
    assert "a<b>" not in missing_chip("a<b>")


# --- verifier bar ----------------------------------------------------------

def test_verifier_summary_prefers_trace_count():
    s = verifier_summary(_answer())
    assert s["verified_n"] == 1
    assert s["dropped_n"] == 1 and s["dropped_source"] == "trace"
    assert any("c2" in r and "unsupported" in r for r in s["dropped_reasons"])
    assert any("not verbatim" in r for r in s["dropped_reasons"])
    assert "trace has no per-claim reasons" in s["reasons_source"]
    assert s["confidence"] == 0.82


def test_verifier_summary_falls_back_to_claims():
    s = verifier_summary(_answer(trace={}))
    assert s["dropped_n"] == 1 and s["dropped_source"] == "claims"


def test_verifier_bar_html_numbers():
    bar = verifier_bar_html(verifier_summary(_answer()))
    assert "1</strong> verified" in bar
    assert "1</strong> dropped" in bar
    assert "0.82" in bar
    assert "not provided" in verifier_bar_html(verifier_summary(_answer(confidence=None)))


# --- contradictions --------------------------------------------------------

def test_diff_marks_only_differences():
    ha, hb = diff_sides("custody of ten days granted", "custody of twelve days granted")
    assert "<mark>ten</mark>" in ha and "<mark>twelve</mark>" in hb
    assert "custody" in ha and "<mark>custody</mark>" not in ha


def test_diff_escapes_html():
    ha, _ = diff_sides("a <b> tag", "a <c> tag")
    assert "<b>" not in ha and "&lt;b&gt;" in ha


def test_contradiction_view_model_fetch_and_failure():
    ans = _answer()
    contra = ans.contradictions[0]
    vm = contradiction_view_model(
        contra, {contra.claim_a: _Chunk("ten days"), contra.claim_b: _Chunk("twelve days")}
    )
    assert vm["text_a"] == "ten days" and vm["text_b"] == "twelve days"
    vm_fail = contradiction_view_model(contra, {contra.claim_a: None, contra.claim_b: None})
    assert vm_fail["text_a"] is None and vm_fail["text_b"] is None
    assert vm_fail["description"] == "custody length differs"


# --- DOCX export -----------------------------------------------------------

def _docs() -> dict:
    return docs_by_id([
        Doc(doc_id="bnss_2023", title="BNSS excerpt", doc_type="statute",
            citation="BNSS, 2023"),
        Doc(doc_id="case_file_demo", title="Demo case file", doc_type="case_file"),
    ])


def test_draft_segments_kinds():
    segs = draft_segments(_answer(), {"FIR number": "0123/2024"})
    kinds = [k for para in segs["body"] for _, k in para]
    assert "marker" in kinds and "text" in kinds
    assert any(t == "[c1]" for para in segs["body"] for t, k in para if k == "marker")
    still_kinds = [[k for _, k in para] for para in segs["still_needed"]]
    assert still_kinds[0] == ["missing", "text", "user"]
    assert still_kinds[1] == ["missing"]


def _docx_texts(blob: bytes) -> tuple[list[str], list[str]]:
    from docx import Document
    from io import BytesIO

    doc = Document(BytesIO(blob))
    paras = [p.text for p in doc.paragraphs]
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    return paras, heads


def test_export_docx_screen_parity_no_extra_text():
    blob = export_draft_docx(_answer(), "Draft a bail prayer?", {"FIR number": "0123/2024"}, _docs())
    paras, heads = _docx_texts(blob)
    assert heads == ["Draft", "Still needed", "Sources"]
    body = "\n".join(paras)
    for needle in ("Draft a bail prayer?", "[c1]", "[MISSING: FIR number]",
                   "[USER-PROVIDED: 0123/2024]", "[MISSING: bail provision]",
                   "bnss_2023::p1::c1", "BNSS excerpt", "BNSS, 2023",
                   "judicial discretion"):
        assert needle in body, needle
    for banned in ("Precheck", "confidence", "Trace", "latency", "retrieved_chunk_ids"):
        assert banned not in body, banned


def test_export_docx_flags_unrenderable_claim():
    ans = _answer(claims=[{
        "claim_id": "c9", "text": "Unsourced assertion.",
        "chunk_ids": [], "quote": "", "status": "unsupported",
    }])
    _, heads = _docx_texts(export_draft_docx(ans, "", {}, _docs()))
    assert heads == ["Draft", "Still needed", "Sources"]
    paras, _ = _docx_texts(export_draft_docx(ans, "", {}, _docs()))
    assert any("no verified quote" in p for p in paras)
